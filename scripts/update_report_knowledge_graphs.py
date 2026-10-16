"""Refresh the three complete skill graphs and code lookup tables in the report."""

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs/bao_cao/Bao_Cao_Thực Tập.docx"
SPECS = (
    {
        "language": "Python", "slug": "python", "figure": "2.5", "table": "2.2",
        "graph": ROOT / "backend/src/infrastructure/data/pythonSkillGraph.json",
    },
    {
        "language": "JavaScript", "slug": "javascript", "figure": "2.6", "table": "2.3",
        "graph": ROOT / "ai-service/data/javascriptSkillGraph.json",
    },
    {
        "language": "C++", "slug": "cpp", "figure": "2.7", "table": "2.4",
        "graph": ROOT / "curriculum/cppSkillGraph.json",
    },
)


def paragraph_before(anchor, text="", style="Normal", bold_prefix=None):
    paragraph = anchor.insert_paragraph_before(style=style)
    if bold_prefix:
        paragraph.add_run(bold_prefix).bold = True
    paragraph.add_run(text)
    return paragraph


def remove_existing_section(doc):
    heading = next((p for p in doc.paragraphs
                    if p.text == "2.3.2. Đồ thị tri thức của ba ngôn ngữ"
                    and p.style.name == "Heading 3"), None)
    if heading is not None:
        tail = next(p for p in doc.paragraphs
                    if p.text.startswith("2.4. Các chức năng quản trị")
                    and p.style.name == "Heading 2")
        current = heading._element
        while current is not tail._element:
            following = current.getnext()
            current.getparent().remove(current)
            current = following
    for paragraph in list(doc.paragraphs):
        if paragraph.text.startswith((
            "Hình 2.5.", "Hình 2.6.", "Hình 2.7.",
            "Bảng 2.2. Tra cứu mã kỹ năng Python",
            "Bảng 2.3. Tra cứu mã kỹ năng JavaScript",
            "Bảng 2.4. Tra cứu mã kỹ năng C++",
        )):
            paragraph._element.getparent().remove(paragraph._element)


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        properties.append(element)
    element.set(qn("w:fill"), fill)


def format_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(1.5), Inches(3.05), Inches(1.65))
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            cell.width = widths[cell_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cell, "25466A" if row_index == 0 else ("F1F6FA" if row_index % 2 == 0 else "FFFFFF"))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.08
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if cell_index != 1 else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9)
                    run.font.bold = row_index == 0
                    if row_index == 0:
                        run.font.color.rgb = RGBColor(255, 255, 255)
    header = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header.append(repeat)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    table._tbl.tblPr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), "85" if side in ("top", "bottom") else "100")
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    table._tbl.tblPr.append(margins)


def add_figure(doc, anchor, spec, graph):
    image = ROOT / f"docs/bao_cao/do_thi_tri_thuc_{spec['slug']}.png"
    picture = anchor.insert_paragraph_before(style="Normal")
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    picture.add_run().add_picture(str(image), width=Inches(6.05))
    caption = (
        f"Hình {spec['figure']}. Đồ thị đầy đủ {len(graph['skills'])} kỹ năng "
        f"và {len(graph['edges'])} quan hệ tiên quyết {spec['language']}"
    )
    caption_paragraph = paragraph_before(anchor, caption, style="Caption1")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return caption


def add_lookup_table(doc, anchor, spec, graph):
    title = f"Bảng {spec['table']}. Tra cứu mã kỹ năng {spec['language']}"
    caption = paragraph_before(anchor, title, style="Caption1")
    caption.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Mã kỹ năng"
    table.rows[0].cells[1].text = "Tên kỹ năng"
    table.rows[0].cells[2].text = "Tiên quyết trực tiếp"
    for skill in graph["skills"]:
        cells = table.add_row().cells
        cells[0].text = skill["id"]
        suffix = " (dự kiến)" if skill.get("content_status") == "PLANNED" else ""
        cells[1].text = skill["name"] + suffix
        cells[2].text = ", ".join(skill.get("prerequisites", [])) or "Không có"
    format_table(table)
    anchor._p.addprevious(table._tbl)
    return title


def main():
    doc = Document(REPORT)
    remove_existing_section(doc)
    for paragraph in doc.paragraphs:
        if paragraph.text == "Bảng 2.2. Các nhóm chức năng đã triển khai":
            for run in paragraph.runs:
                run.text = run.text.replace("Bảng 2.2.", "Bảng 2.5.")
    anchor = next(p for p in doc.paragraphs
                  if p.text.startswith("2.4. Các chức năng quản trị")
                  and p.style.name == "Heading 2")
    graphs = {spec["slug"]: json.loads(spec["graph"].read_text(encoding="utf-8")) for spec in SPECS}

    paragraph_before(anchor, "2.3.2. Đồ thị tri thức của ba ngôn ngữ", style="Heading 3")
    paragraph_before(
        anchor,
        "Hệ thống dùng đồ thị có hướng không chu trình (DAG) để mô tả điều kiện tiên quyết giữa các kỹ năng. "
        "Mũi tên A → B nghĩa là cần nắm A trước khi học B; một kỹ năng có thể có nhiều tiền đề. "
        "Module và topic chỉ phân nhóm nội dung, không phải cạnh tiên quyết. Các hình dưới đây vẽ tất cả kỹ năng và "
        "quan hệ trong từng tệp đồ thị; mỗi ô chỉ ghi mã để hình không quá dày. Bảng ngay sau mỗi hình giải nghĩa toàn bộ mã "
        "và liệt kê các tiền đề trực tiếp.",
    )
    paragraph_before(
        anchor,
        "Phiên bản vận hành hiện tại gồm Python v2.1 (35 kỹ năng, 49 cạnh, 50 ánh xạ bài học), "
        "JavaScript v3.1.0 (23 kỹ năng, 29 cạnh, 60 ánh xạ) và C++ v4.1.0 (21 kỹ năng, 22 cạnh, 48 ánh xạ). "
        "Màu xám trong hình Python chỉ kỹ năng có trạng thái PLANNED; màu xanh chỉ các kỹ năng còn lại. "
        "Số lượng nút/cạnh trong chú thích hình được đếm từ JSON, không phải ước lượng từ một nhánh mẫu.",
    )

    summaries = {
        "python": (
            "Từ biến, toán tử và điều kiện, các nhánh phát triển sang vòng lặp, cấu trúc dữ liệu, hàm, "
            "ngoại lệ, module và hướng đối tượng. Có 8 kỹ năng PLANNED nên cần kiểm tra trạng thái học liệu "
            "khi dùng đồ thị để chọn bài học."
        ),
        "javascript": (
            "Từ biến và kiểu dữ liệu, đồ thị mở rộng sang nhập liệu runtime, hàm, mảng, đối tượng và các "
            "chủ đề nâng cao. Execution Context là tiền đề trực tiếp của Scope, rồi Scope dẫn tới Closure, "
            "phù hợp với thứ tự hai bài 11.01 và 11.02 trong giáo trình."
        ),
        "cpp": (
            "Nhánh điều khiển đi từ điều kiện qua vòng lặp, mảng tới ma trận. Nhánh con trỏ nhận hai tiền đề "
            "trực tiếp là struct và vòng đời đối tượng; từ con trỏ mới tới RAII/smart pointer. "
            "Các cạnh bắc cầu dư đã được bỏ khỏi bản vận hành."
        ),
    }
    figure_entries = []
    table_entries = []
    for spec in SPECS:
        graph = graphs[spec["slug"]]
        paragraph_before(anchor, spec["language"], style="Heading 4")
        paragraph_before(anchor, summaries[spec["slug"]])
        figure_entries.append(add_figure(doc, anchor, spec, graph))
        table_entries.append(add_lookup_table(doc, anchor, spec, graph))

    paragraph_before(
        anchor,
        "Các quan hệ trên là thiết kế chương trình học, chưa chứng minh thứ tự này tối ưu cho mọi người học. "
        "Kết quả BKT, DKT và PAL-Net ở Chương 4–5 vẫn dựa trên snapshot JavaScript v3.0.0, C++ v4.0.0 "
        "và Python v2.1; chưa có thí nghiệm huấn luyện lại với hai đồ thị vận hành mới.",
    )

    figure_list_anchor = next(p for p in doc.paragraphs if p.text.startswith("Hình 3.1."))
    for entry in figure_entries:
        paragraph_before(figure_list_anchor, entry)
    table_list_anchor = next(p for p in doc.paragraphs
                             if p.text == "Bảng 2.5. Các nhóm chức năng đã triển khai"
                             and p.style.name != "Caption1")
    for entry in table_entries:
        paragraph_before(table_list_anchor, entry)

    generated = REPORT.with_name("Bao_Cao_Thuc_Tap.generated.docx")
    doc.save(generated)
    Document(generated)
    generated.replace(REPORT)
    print("Đã cập nhật ba đồ thị đầy đủ và ba bảng tra cứu trong báo cáo Word.")


if __name__ == "__main__":
    main()
