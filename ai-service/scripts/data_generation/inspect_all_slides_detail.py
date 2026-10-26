import sys
from pathlib import Path
from pptx import Presentation

if sys.stdout.encoding != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")

p = Path(r"D:\Project\LearnPython\docs\slide-thuyet_trinh\Blue Professional Data Analytics Presentation.pptx")
prs = Presentation(str(p))

for i, slide in enumerate(prs.slides):
    print(f"=== SLIDE {i+1} ===")
    for sh_idx, shape in enumerate(slide.shapes):
        txt = ""
        if shape.has_text_frame:
            txt = " /// ".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
        print(f"  Shape {sh_idx} [{shape.name}, type={shape.shape_type}]: {txt[:70]}")
