import sys
from pathlib import Path
from pptx import Presentation

if sys.stdout.encoding != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")

p = Path(r"D:\Project\LearnPython\docs\slide-thuyet_trinh\Blue Professional Data Analytics Presentation.pptx")
prs = Presentation(str(p))

for i in [0, 1, 2, 3, 5, 10]:
    s = prs.slides[i]
    print(f"=== SLIDE {i+1} ===")
    for shape in s.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                if para.text.strip():
                    runs_info = []
                    for r in para.runs:
                        color = "None"
                        try:
                            if r.font.color and r.font.color.type == 1:
                                color = str(r.font.color.rgb)
                            elif r.font.color:
                                color = f"Theme:{r.font.color.theme_color}"
                        except Exception as e:
                            color = str(e)
                        fsize = f"{r.font.size.pt:.1f}pt" if r.font.size else "inherit"
                        runs_info.append(f"'{r.text}' (f={r.font.name}, sz={fsize}, b={r.font.bold}, c={color})")
                    if runs_info:
                        print("  Para:", " | ".join(runs_info))
                    else:
                        print("  Para text:", para.text.strip()[:60])
        elif shape.shape_type == 1: # AutoShape
            fill = "None"
            try:
                if shape.fill and shape.fill.type == 1:
                    fill = str(shape.fill.fore_color.rgb)
            except Exception:
                pass
            print(f"  Shape {shape.name}: pos=({shape.left.inches:.2f}, {shape.top.inches:.2f}, {shape.width.inches:.2f}, {shape.height.inches:.2f}) fill={fill}")
    print()
