import sys
from pathlib import Path
from pptx import Presentation

if sys.stdout.encoding != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")

p = Path(r"D:\Project\LearnPython\docs\slide-thuyet_trinh\Blue Professional Data Analytics Presentation.pptx")
prs = Presentation(str(p))

for i in range(len(prs.slides)):
    s = prs.slides[i]
    shapes_summary = []
    for sh in s.shapes:
        name = sh.name
        stype = sh.shape_type
        pos = f"({sh.left.inches:.1f}, {sh.top.inches:.1f}, {sh.width.inches:.1f}, {sh.height.inches:.1f})"
        text = ""
        if sh.has_text_frame and sh.text_frame.text.strip():
            text = " | ".join(p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip())[:50]
        shapes_summary.append(f"[{name}: type={stype}, pos={pos}, text='{text}']")
    print(f"Slide {i+1} ({len(s.shapes)} shapes):")
    for sm in shapes_summary:
        print("  ", sm)
    print()
