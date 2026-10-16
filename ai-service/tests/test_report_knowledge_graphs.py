import hashlib
import json
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "docs/bao_cao/Bao_Cao_Thực Tập.docx"
SPECS = (
    ("python", "2.5", ROOT / "backend/src/infrastructure/data/pythonSkillGraph.json"),
    ("javascript", "2.6", ROOT / "ai-service/data/javascriptSkillGraph.json"),
    ("cpp", "2.7", ROOT / "curriculum/cppSkillGraph.json"),
)


class ReportKnowledgeGraphTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.doc = Document(REPORT)
        cls.lookup_tables = [
            table for table in cls.doc.tables
            if table.cell(0, 0).text == "Mã kỹ năng"
        ]

    def test_every_skill_is_explained_once_with_direct_prerequisites(self):
        self.assertEqual(len(self.lookup_tables), 3)
        for (_, _, graph_path), table in zip(SPECS, self.lookup_tables):
            graph = json.loads(graph_path.read_text(encoding="utf-8"))
            skills = {skill["id"]: skill for skill in graph["skills"]}
            rows = {
                row.cells[0].text: (row.cells[1].text, row.cells[2].text)
                for row in table.rows[1:]
            }
            self.assertEqual(len(rows), len(table.rows) - 1)
            self.assertEqual(set(rows), set(skills))
            for skill_id, skill in skills.items():
                name, prerequisites = rows[skill_id]
                self.assertTrue(name.startswith(skill["name"]))
                self.assertEqual(
                    prerequisites,
                    ", ".join(skill["prerequisites"]) or "Không có",
                )

    def test_full_graph_images_match_the_embedded_artwork(self):
        paragraphs = self.doc.paragraphs
        for language, number, graph_path in SPECS:
            graph = json.loads(graph_path.read_text(encoding="utf-8"))
            caption_index = next(
                index for index, paragraph in enumerate(paragraphs)
                if paragraph.style.name == "Caption1"
                and paragraph.text.startswith(f"Hình {number}.")
            )
            caption = paragraphs[caption_index].text
            self.assertIn(f"{len(graph['skills'])} kỹ năng", caption)
            self.assertIn(f"{len(graph['edges'])} quan hệ", caption)
            blip = paragraphs[caption_index - 1]._element.xpath(".//a:blip")[0]
            embedded = self.doc.part.related_parts[blip.get(qn("r:embed"))].blob
            source = ROOT / f"docs/bao_cao/do_thi_tri_thuc_{language}.png"
            self.assertEqual(
                hashlib.sha256(embedded).digest(),
                hashlib.sha256(source.read_bytes()).digest(),
            )


if __name__ == "__main__":
    unittest.main()
